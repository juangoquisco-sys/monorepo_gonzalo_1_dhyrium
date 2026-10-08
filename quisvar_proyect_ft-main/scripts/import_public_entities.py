"""Importa todas las hojas del directorio público a un catálogo web verificable."""

from __future__ import annotations

import json
import re
import unicodedata
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook


SOURCE_DIRECTORY = Path(
    r"C:\Users\unidad lineal\Documents\Codex\2026-09-21\referenced-chatgpt-conversation-this-is-an\outputs\directorio-entidades-publicas-peru"
)
SOURCE_FILE = SOURCE_DIRECTORY / "directorio_entidades_publicas_peru_2026_con_codigos_ue.xlsx"
INSPECTION_FILE = SOURCE_DIRECTORY / "directorio_entidades_publicas_peru_2026_con_codigos_ue.xlsx.inspect.ndjson"
PROJECT_DIRECTORY = Path(__file__).resolve().parent.parent
OUTPUT_FILE = PROJECT_DIRECTORY / "public" / "data" / "entidades-publicas-peru-2026.json"

DATASETS = {
    "entities": {
        "sheet": "Entidades públicas",
        "prefix": "entidad-publica",
        "requiredIndex": 3,
        "fields": [
            "departamento", "provincia", "distrito", "entidad", "ruc", "codigoUeSiaf", "ubigeo",
            "titular", "cargo", "direccion", "telefonoOCelular", "anexo", "paginaWebOficial",
            "fuenteWeb", "fuenteContacto", "fuenteDatos", "nivelDeGobierno", "codigoPliego",
            "pliego", "nombreMef", "validacionCodigo", "fuenteCodigoMef",
        ],
    },
    "municipalities": {
        "sheet": "Municipalidades",
        "prefix": "municipalidad",
        "requiredIndex": 2,
        "fields": [
            "departamento", "provincia", "distrito", "alcalde", "cargo", "direccion", "codigoCiudad",
            "telefono", "anexo", "correoElectronico", "ubigeo", "codigoSiaf", "paginaWebOficial",
            "fuenteWeb", "fuente", "validacionCodigo", "fuenteCodigoMef",
        ],
    },
    "mefCatalog": {
        "sheet": "Catálogo MEF 2026",
        "prefix": "catalogo-mef",
        "requiredIndex": 0,
        "fields": [
            "codigoUeSiaf", "nombreMef", "nivelDeGobierno", "codigoPliego", "pliego",
            "departamento", "provincia", "distrito", "ubigeo", "fuenteMef",
        ],
    },
}


def to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def import_dataset(workbook, definition: dict) -> tuple[list[dict[str, str | int]], list[str]]:
    worksheet = workbook[definition["sheet"]]
    fields = definition["fields"]
    headers = [to_text(worksheet.cell(4, column).value) for column in range(1, len(fields) + 1)]
    records: list[dict[str, str | int]] = []

    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=5, max_col=len(fields), values_only=True),
        start=5,
    ):
        values = [to_text(value) for value in row]
        if not any(values) or not values[definition["requiredIndex"]]:
            continue
        records.append({
            "id": f'{definition["prefix"]}-{row_number}',
            "sourceRow": row_number,
            **dict(zip(fields, values)),
        })

    return records, headers


def read_inspection_summary() -> list[dict]:
    summary: list[dict] = []
    for line in INSPECTION_FILE.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        entry = json.loads(line)
        if entry.get("kind") in {"workbook", "sheet"}:
            summary.append(entry)
        elif entry.get("kind") == "table":
            summary.append({key: entry.get(key) for key in ("kind", "sheet", "address", "rows", "cols")})
    return summary


UNIFIED_FIELDS = [
    "nombrePrincipal", "ruc", "codigoUeSiaf", "ubigeo", "departamento", "provincia", "distrito",
    "nivelDeGobierno", "codigoPliego", "pliego", "nombreMef", "titular", "cargo", "alcalde",
    "cargoAlcalde", "direccion", "direccionMunicipal", "telefonoOCelular", "codigoCiudad",
    "telefonoMunicipal", "anexo", "anexoMunicipal", "correoElectronico", "paginaWebOficial",
    "paginaWebMunicipal", "fuenteWeb", "fuenteWebMunicipal", "fuenteContacto", "fuenteDatos",
    "fuenteMunicipal", "fuenteMef", "validacionCodigo", "validacionCodigoMunicipal",
    "fuenteCodigoMef", "fuenteCodigoMefMunicipal",
]


def normalize_identifier(value: object, compact: bool = False) -> str:
    normalized = unicodedata.normalize("NFD", to_text(value).upper()).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", "" if compact else " ", normalized).strip()


def create_unified(record_id: str) -> dict:
    return {
        "id": record_id,
        **{field: "" for field in UNIFIED_FIELDS},
        "aliases": [],
        "sources": [],
        "sourceRows": {"entities": [], "municipalities": [], "mefCatalog": []},
        "matchMethods": [],
    }


def set_if_empty(target: dict, field: str, value: object) -> None:
    text = to_text(value)
    if text and not target.get(field):
        target[field] = text


def append_unique(target: list, value: object) -> None:
    text = to_text(value)
    if text and text not in target:
        target.append(text)


def add_source(target: dict, dataset_key: str, source_label: str, source_row: int, method: str) -> None:
    append_unique(target["sources"], source_label)
    if source_row not in target["sourceRows"][dataset_key]:
        target["sourceRows"][dataset_key].append(source_row)
    append_unique(target["matchMethods"], method)


def merge_mef(target: dict, source: dict, method: str = "Código UE/SIAF") -> None:
    append_unique(target["aliases"], source.get("nombreMef"))
    for field in ("codigoUeSiaf", "nombreMef", "nivelDeGobierno", "codigoPliego", "pliego",
                  "departamento", "provincia", "distrito", "ubigeo", "fuenteMef"):
        set_if_empty(target, field, source.get(field))
    set_if_empty(target, "nombrePrincipal", source.get("nombreMef"))
    add_source(target, "mefCatalog", "Catálogo MEF 2026", int(source["sourceRow"]), method)


def merge_municipality(target: dict, source: dict, method: str = "Código SIAF") -> None:
    municipality_label = f"Municipalidad de {source.get('distrito') or source.get('provincia') or source.get('departamento')}"
    append_unique(target["aliases"], municipality_label)
    common_fields = ("departamento", "provincia", "distrito", "ubigeo")
    for field in common_fields:
        set_if_empty(target, field, source.get(field))
    set_if_empty(target, "codigoUeSiaf", source.get("codigoSiaf"))
    mapping = {
        "alcalde": "alcalde", "cargoAlcalde": "cargo", "direccionMunicipal": "direccion",
        "codigoCiudad": "codigoCiudad", "telefonoMunicipal": "telefono", "anexoMunicipal": "anexo",
        "correoElectronico": "correoElectronico", "paginaWebMunicipal": "paginaWebOficial",
        "fuenteWebMunicipal": "fuenteWeb", "fuenteMunicipal": "fuente",
        "validacionCodigoMunicipal": "validacionCodigo", "fuenteCodigoMefMunicipal": "fuenteCodigoMef",
    }
    for target_field, source_field in mapping.items():
        set_if_empty(target, target_field, source.get(source_field))
    set_if_empty(target, "nombrePrincipal", municipality_label)
    add_source(target, "municipalities", "Municipalidades", int(source["sourceRow"]), method)


def merge_entity(target: dict, source: dict, method: str) -> None:
    previous_name = target.get("nombrePrincipal")
    append_unique(target["aliases"], previous_name)
    append_unique(target["aliases"], source.get("entidad"))
    source_name = to_text(source.get("entidad"))
    if source_name:
        target["nombrePrincipal"] = source_name
    for field in ("ruc", "codigoUeSiaf", "ubigeo", "departamento", "provincia", "distrito",
                  "nivelDeGobierno", "codigoPliego", "pliego", "nombreMef", "titular", "cargo",
                  "direccion", "telefonoOCelular", "anexo", "paginaWebOficial", "fuenteWeb",
                  "fuenteContacto", "fuenteDatos", "validacionCodigo", "fuenteCodigoMef"):
        set_if_empty(target, field, source.get(field))
    add_source(target, "entities", "Entidades públicas", int(source["sourceRow"]), method)


def municipality_name_matches(entity: dict, municipality_record: dict) -> bool:
    entity_name = normalize_identifier(entity.get("entidad"), compact=True)
    if "MUNICIPALIDADDISTRITAL" in entity_name:
        place = normalize_identifier(municipality_record.get("distrito"), compact=True)
        entity_place = entity_name.replace("MUNICIPALIDADDISTRITALDE", "").replace("MUNICIPALIDADDISTRITALLA", "")
    elif "MUNICIPALIDADPROVINCIAL" in entity_name:
        place = normalize_identifier(municipality_record.get("provincia"), compact=True)
        entity_place = entity_name.replace("MUNICIPALIDADPROVINCIALDE", "").replace("MUNICIPALIDADPROVINCIALDEL", "")
    else:
        return False
    return bool(place) and (place in entity_name or SequenceMatcher(None, place, entity_place).ratio() >= 0.72)


def normalized_match_name(value: object) -> str:
    stop_words = {"DE", "DEL", "LA", "LAS", "LOS", "Y", "EL"}
    return " ".join(
        word for word in normalize_identifier(value).split()
        if word not in stop_words
    )


def high_confidence_mef_match(source: dict, candidates: list[dict]) -> dict | None:
    source_name = normalized_match_name(source.get("entidad"))
    if not source_name or not candidates:
        return None

    scored = sorted(
        (
            (
                SequenceMatcher(None, source_name, normalized_match_name(candidate.get("nombreMef"))).ratio(),
                candidate,
            )
            for candidate in candidates
        ),
        key=lambda item: item[0],
    )
    best_score, best_candidate = scored[-1]
    second_score = scored[-2][0] if len(scored) > 1 else 0
    if best_score >= 0.90 and best_score - second_score >= 0.08:
        return best_candidate
    return None


def unify_datasets(imported: dict[str, list[dict[str, str | int]]]) -> list[dict]:
    unified: list[dict] = []
    by_code: dict[str, dict] = {}
    by_mef_name: dict[str, list[dict]] = defaultdict(list)
    by_mef_ubigeo: dict[str, list[dict]] = defaultdict(list)

    for source in imported["mefCatalog"]:
        code = to_text(source.get("codigoUeSiaf"))
        target = create_unified(f"entidad-unificada-{code}")
        merge_mef(target, source, "Registro base MEF")
        unified.append(target)
        by_code[code] = target
        by_mef_name[normalize_identifier(source.get("nombreMef"))].append(target)
        ubigeo = to_text(source.get("ubigeo"))
        if ubigeo:
            by_mef_ubigeo[ubigeo].append(target)

    municipality_by_ubigeo: dict[str, list[dict]] = defaultdict(list)
    for source in imported["municipalities"]:
        code = to_text(source.get("codigoSiaf"))
        target = by_code.get(code) if code else None
        if target is None:
            target = create_unified(f"entidad-unificada-municipalidad-{source['sourceRow']}")
            unified.append(target)
        merge_municipality(target, source, "Código SIAF" if code else "Municipalidad sin código SIAF")
        ubigeo = to_text(source.get("ubigeo"))
        if ubigeo and target not in municipality_by_ubigeo[ubigeo]:
            municipality_by_ubigeo[ubigeo].append(target)

    by_pte_identity: dict[tuple[str, str], dict] = {}
    for source in imported["entities"]:
        code = to_text(source.get("codigoUeSiaf"))
        normalized_name = normalize_identifier(source.get("entidad"))
        ruc = to_text(source.get("ruc"))
        identity = (ruc, normalized_name)
        target = by_code.get(code) if code else None
        method = "Código UE/SIAF"

        if target is None and normalized_name and len(by_mef_name[normalized_name]) == 1:
            target = by_mef_name[normalized_name][0]
            method = "Nombre MEF exacto"

        if target is None:
            target = high_confidence_mef_match(
                source,
                by_mef_ubigeo.get(to_text(source.get("ubigeo")), []),
            )
            if target is not None:
                method = "Nombre y UBIGEO de alta confianza"

        if target is None:
            ubigeo_matches = municipality_by_ubigeo.get(to_text(source.get("ubigeo")), [])
            if len(ubigeo_matches) == 1 and municipality_name_matches(source, ubigeo_matches[0]):
                target = ubigeo_matches[0]
                method = "Municipalidad y UBIGEO verificados"

        if target is None and ruc and normalized_name and identity in by_pte_identity:
            target = by_pte_identity[identity]
            method = "RUC y nombre exactos"

        if target is None:
            target = create_unified(f"entidad-unificada-pte-{source['sourceRow']}")
            unified.append(target)
            method = "Registro PTE sin coincidencia segura"

        merge_entity(target, source, method)
        if ruc and normalized_name:
            by_pte_identity.setdefault(identity, target)

    for target in unified:
        target["aliases"] = [alias for alias in target["aliases"] if alias != target["nombrePrincipal"]]
        target["sourceReferenceCount"] = sum(len(rows) for rows in target["sourceRows"].values())

    return unified


workbook = load_workbook(SOURCE_FILE, read_only=True, data_only=True)
imported: dict[str, list[dict[str, str | int]]] = {}
dataset_metadata: dict[str, dict] = {}

for dataset_key, definition in DATASETS.items():
    records, headers = import_dataset(workbook, definition)
    imported[dataset_key] = records
    dataset_metadata[dataset_key] = {
        "sourceSheet": definition["sheet"],
        "recordCount": len(records),
        "sourceHeaders": headers,
    }

overview_sheet = workbook["Por departamento"]
overview_rows = [
    {"sourceRow": row_number, "values": [to_text(value) for value in row]}
    for row_number, row in enumerate(overview_sheet.iter_rows(max_col=4, values_only=True), start=1)
    if any(to_text(value) for value in row)
]

total_records = sum(len(records) for records in imported.values())
unified_entities = unify_datasets(imported)
payload = {
    "metadata": {
        "title": "Directorio de entidades públicas del Perú 2026",
        "sourceFile": SOURCE_FILE.name,
        "sourceNotesFile": INSPECTION_FILE.name,
        "headerRow": 4,
        "importedRecords": total_records,
        "datasets": dataset_metadata,
        "unification": {
            "sourceRecords": total_records,
            "unifiedRecords": len(unified_entities),
            "consolidatedReferences": total_records - len(unified_entities),
            "primaryKey": "Código UE/SIAF",
        },
    },
    **imported,
    "unifiedEntities": unified_entities,
    "departmentOverview": {
        "sourceSheet": overview_sheet.title,
        "rows": overview_rows,
    },
    "inspectionSummary": read_inspection_summary(),
}

OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(
    f"Importados {total_records} registros: "
    f"{len(imported['entities'])} entidades, "
    f"{len(imported['municipalities'])} municipalidades y "
    f"{len(imported['mefCatalog'])} registros MEF. "
    f"Directorio unificado: {len(unified_entities)} entidades."
)
